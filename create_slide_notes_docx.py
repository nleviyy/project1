# -*- coding: utf-8 -*-
"""
Script tạo file Word (.docx) chứa Slide & Kịch bản thuyết trình chi tiết
Dành cho: Nguyễn Lê Vy (Team Leader) và nhóm đồ án F&B
"""
import os
import sys
from docx import Document
if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'))

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def generate_slide_docx():
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    COLOR_PRIMARY = RGBColor(15, 23, 42)     # Slate 900
    COLOR_BLUE = RGBColor(37, 99, 235)       # Royal Blue
    COLOR_SECONDARY = RGBColor(71, 85, 105)  # Slate 600
    COLOR_TEXT = RGBColor(30, 41, 59)        # Slate 800

    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = COLOR_TEXT
    normal_style.paragraph_format.line_spacing = 1.2
    normal_style.paragraph_format.space_after = Pt(4)

    # HEADER
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_inst = p_inst.add_run("BÁO CÁO TIẾN ĐỘ & NGHIỆM THU ĐỒ ÁN MODULE 1 (PYTHON)\nBỘ MÔN CÔNG NGHỆ THÔNG TIN")
    r_inst.font.size = Pt(9.5)
    r_inst.font.bold = True
    r_inst.font.color.rgb = COLOR_SECONDARY

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(8)
    p_title.paragraph_format.space_after = Pt(4)
    r_title = p_title.add_run("NỘI DUNG SLIDE & KỊCH BẢN THUYẾT TRÌNH BÁO CÁO")
    r_title.font.size = Pt(16)
    r_title.font.bold = True
    r_title.font.color.rgb = COLOR_BLUE

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(14)
    r_sub = p_sub.add_run("Hệ Thống Quản Lý Đơn Hàng & Bàn Ăn Nhà Hàng F&B (Nhánh dev - 100% Hoàn Tất)")
    r_sub.font.size = Pt(11.5)
    r_sub.font.italic = True
    r_sub.font.color.rgb = COLOR_PRIMARY

    # META TABLE
    t_meta = doc.add_table(rows=4, cols=2)
    t_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta = [
        ("Nhóm sinh viên thực hiện:", "Nguyễn Lê Vy (Leader), Lê Thị Như Quỳnh, Ngô Thị Lan Anh"),
        ("Mã nguồn GitHub (Nhánh dev):", "https://github.com/nleviyy/project1 (Đã merge 100%, 125 test pass)"),
        ("Thời lượng thuyết trình dự kiến:", "10 - 15 phút (Kèm kịch bản lời thoại chi tiết từng slide)"),
        ("Thời điểm báo cáo:", "Đánh giá tiến độ hoàn thiện & Kế hoạch nghiệm thu đồ án"),
    ]
    for idx, (k, v) in enumerate(meta):
        row = t_meta.rows[idx]
        c1, c2 = row.cells[0], row.cells[1]
        c1.text = k
        c2.text = v
        c1.paragraphs[0].runs[0].font.bold = True
        c1.paragraphs[0].runs[0].font.color.rgb = COLOR_BLUE
        c1.width = Inches(2.3)
        c2.width = Inches(4.5)
        set_cell_background(c1, "F1F5F9")
        set_cell_background(c2, "F8FAFC")
        set_cell_margins(c1, 50, 50, 80, 80)
        set_cell_margins(c2, 50, 50, 80, 80)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    slides_data = [
        (
            "Slide 1: Trang Tiêu Đề & Giới Thiệu Đề Tài",
            [
                "Tên đề tài: Hệ Thống Quản Lý Nhà Hàng F&B Bằng Python & Streamlit",
                "Thành viên & Phân công: Nguyễn Lê Vy (Team Leader/Architecture/Tables/UI), Lê Thị Như Quỳnh (Cart/Payments/Invoices), Ngô Thị Lan Anh (Kitchen FIFO/Data Storage/Structures)",
                "Mã nguồn: GitHub nleviyy/project1 - Nhánh chính: dev (Đã merge và test hoàn tất)",
            ],
            "Kính chào Thầy/Cô và các bạn! Hôm nay, nhóm chúng em xin phép trình bày báo cáo tiến độ và kết quả hoàn thiện hệ thống Quản lý Nhà hàng F&B. Dự án được xây dựng dựa trên sự tiếp thu toàn diện các ý kiến đóng góp từ buổi đánh giá trước của Mentor. Toàn bộ mã nguồn đã được chuẩn hóa, hoàn thiện luồng nghiệp vụ khép kín và tích hợp đầy đủ trên nhánh dev của GitHub với 125 test case tự động đạt 100%."
        ),
        (
            "Slide 2: Tiếp Thu Feedback Mentor & Mục Tiêu Cải Tiến",
            [
                "Vấn đề tồn tại trước đây: Code phân mảnh trên các nhánh riêng rẽ; giao diện còn phụ thuộc dữ liệu mô phỏng, dùng hàm sleep(), random(); bếp và thu ngân chưa đồng bộ.",
                "Giải pháp thực hiện: Hợp nhất toàn bộ vào cấu trúc 3 tầng chuẩn (models/core/ui); sử dụng 100% dữ liệu thực từ 5 tệp JSON; khép kín luồng từ Đặt món -> Bếp -> Thanh toán -> Trả bàn.",
            ],
            "Ở giai đoạn trước, Mentor đã chỉ ra hạn chế lớn nhất của nhóm là các nhánh tính năng hoạt động rời rạc, giao diện còn dùng dữ liệu giả lập và chưa có sự đồng bộ giữa bếp với thu ngân. Tiếp thu nghiêm túc nhận xét đó, nhóm em đã đặt ra mục tiêu cốt lõi: Hợp nhất toàn bộ mã nguồn về nhánh dev, chuyển đổi 100% sang dữ liệu thực từ các file JSON, và loại bỏ triệt để mọi logic giả lập."
        ),
        (
            "Slide 3: Bảng Phân Công Nhiệm Vụ & Tiến Độ (Google Sheets)",
            [
                "Task 1 (Vy): Tổ chức codebase 3 tầng, chỉ dùng src/app.py làm entry point duy nhất (100%).",
                "Task 2 (Quỳnh): src/core/Cart.py, kết nối menu.json, thêm tính năng Undo/Redo bằng Stack (100%).",
                "Task 3 (Vy): Tạo đơn PENDING, kiểm soát chiếm bàn AVAILABLE, chống trùng đơn đang phục vụ (100%).",
                "Task 4 (Lan Anh): KitchenManager hàng đợi FIFO, tính prep_time và ước tính thời gian chờ (100%).",
                "Task 5 (Quỳnh): PaymentService, CASH/CARD/TRANSFER, chống thanh toán trùng, tạo Invoice.txt (100%).",
                "Task 6 (Vy): Tích hợp giao diện Streamlit với Core Service, thiết kế Clean Slate Navy/Blue (100%).",
                "Task 7 (Lan Anh): Chuẩn hóa schema JSON, cài đặt DoublyLinkedList, HashTable, Stack (100%).",
                "Task 8 (Cả nhóm): 125 Pytest tự động, 3 script demo CLI, tài liệu README chuẩn (100%).",
            ],
            "Trên màn hình là bảng phân công nhiệm vụ cụ thể của 3 thành viên theo kế hoạch Google Sheets đã thống nhất. Bạn Vy đảm nhiệm cấu trúc nền tảng, quản lý bàn và tích hợp toàn bộ giao diện; Bạn Quỳnh phụ trách module Giỏ hàng và quy trình Thanh toán - Xuất hóa đơn; Bạn Lan Anh phụ trách Điều phối bếp theo hàng đợi FIFO và Cấu trúc dữ liệu tự cài đặt. Đến thời điểm hiện tại, 100% các đầu việc đều đã được kiểm thử và đẩy thành công lên nhánh dev."
        ),
        (
            "Slide 4: Kiến Trúc Hệ Thống & Tổ Chức Codebase",
            [
                "Models (src/models/): Định nghĩa thực thể bằng Pydantic và Dataclasses (Order, MenuItem, Table, Payment, Invoice).",
                "Core (src/core/): Chứa toàn bộ nghiệp vụ thuần Python (Cart, KitchenManager, OrderService, PaymentService, TableManager) độc lập với UI.",
                "UI (src/ui/ & src/app.py): Entry point Streamlit duy nhất, giao tiếp với Core qua session_state.",
                "Storage (data/): menu.json, tables.json, orders.json, payments.json, invoices.json đảm bảo tính toàn vẹn dữ liệu.",
            ],
            "Về mặt kiến trúc, dự án được tổ chức chặt chẽ theo nguyên lý phân tách trách nhiệm. Tầng Models định nghĩa dữ liệu với validation nghiêm ngặt; Tầng Core chứa toàn bộ logic tính toán tiền bạc, chuyển đổi trạng thái và hàng đợi; Tầng UI chỉ đóng vai trò hiển thị và bắt sự kiện từ người dùng. Điểm đặc biệt là toàn bộ tầng Core có thể chạy độc lập qua CLI hoặc UnitTest mà không cần khởi động Streamlit."
        ),
        (
            "Slide 5: Giỏ Hàng, Thực Đơn & Cơ Chế Hoàn Tác (Undo/Redo)",
            [
                "MenuCatalog: Đọc từ data/menu.json, tìm kiếm autocomplete, lọc danh mục, kiểm tra tình trạng còn món và thời gian chế biến prep_time.",
                "Cart.py: Quản lý giỏ hàng động, cập nhật số lượng, ghi chú món ăn (ví dụ: 'Ít đường', 'Nhiều đá').",
                "Cấu trúc Stack hoàn tác: Lưu trữ snapshot của giỏ hàng sau mỗi thao tác, hỗ trợ Undo và Redo linh hoạt.",
            ],
            "Tại phân hệ Đặt món do bạn Quỳnh phụ trách, giỏ hàng không còn bị hard-code mà liên kết trực tiếp với menu.json. Điểm sáng tạo ở đây là nhóm đã tích hợp cơ chế hoàn tác Undo/Redo sử dụng cấu trúc dữ liệu Stack. Khi nhân viên vô tình bấm nhầm số lượng hoặc xóa nhầm món, hệ thống có thể khôi phục lại trạng thái giỏ hàng trước đó chỉ với 1 cú click."
        ),
        (
            "Slide 6: Tạo Đơn & Quản Lý Bàn Ăn Thông Minh",
            [
                "Ràng buộc nghiệp vụ: Bàn phải ở trạng thái AVAILABLE hoặc chọn Mang về (Takeaway); Giỏ hàng không rỗng.",
                "Chống trùng đơn: Khóa bàn, chuyển sang OCCUPIED ngay khi tạo đơn PENDING; không cho phép đặt thêm đơn mới vào bàn đang có khách.",
                "Quản lý bàn linh hoạt: Hỗ trợ đặt trước bàn (RESERVED), hủy bàn đặt trước, thêm/xóa bàn ăn trực tiếp vào tables.json.",
            ],
            "Tiếp theo là phân hệ Tạo đơn và Quản lý bàn do em trực tiếp phụ trách. Hệ thống áp dụng quy tắc nghiệp vụ nghiêm ngặt: Một bàn đang có đơn hoạt động thì không thể đặt đè; Bàn chỉ được chiếm giữ khi đơn hàng được tạo hợp lệ. Ngay khi bấm 'Xác nhận đơn', mã đơn duy nhất dạng ORD-YYYYMMDD-XXX sẽ được sinh tự động, bàn ăn chuyển sang OCCUPIED và dữ liệu lập tức đồng bộ sang khu vực Bếp."
        ),
        (
            "Slide 7: Điều Phối Bếp FIFO & Ước Tính Thời Gian Chờ",
            [
                "Hàng đợi FIFO: Các đơn hàng PENDING được đưa vào Queue theo thứ tự thời gian đặt trước.",
                "Xử lý đơn tuần tự: Bếp chỉ nấu 1 đơn chính (PREPARING) tại một thời điểm để đảm bảo chất lượng phục vụ.",
                "Thuật toán ước tính thời gian chờ: Thời gian còn lại đơn đang nấu + Tổng prep_time các đơn chờ phía trước + prep_time đơn hiện tại.",
                "Vòng đời chuẩn: PENDING -> PREPARING -> SERVED. Đơn đang nấu không thể bị hủy tùy tiện.",
            ],
            "Phân hệ Bếp do bạn Lan Anh phụ trách giải quyết triệt để bài toán điều phối theo mô hình hàng đợi FIFO (First In First Out). Bếp nhận đơn tuần tự, tự động tính toán thời gian chờ dựa trên tổng thời gian chuẩn bị prep_time của các món. Đơn đang nấu sẽ không thể bị hủy tùy tiện, và khi hoàn tất chế biến, bếp bấm 'Lên món' để chuyển đơn sang SERVED, sẵn sàng cho khâu thanh toán."
        ),
        (
            "Slide 8: Thanh Toán An Toàn & Chống Thanh Toán Trùng",
            [
                "Chỉ thanh toán đơn SERVED: Ngăn chặn thu ngân thanh toán đơn khi bếp chưa hoàn thành món.",
                "Chống thanh toán trùng (Idempotency): Kiểm tra lịch sử giao dịch; từ chối và cảnh báo nếu đơn đã thanh toán thành công.",
                "Đa dạng phương thức: Tiền mặt (tự tính tiền thừa trả khách), Chuyển khoản, Thẻ ngân hàng (lưu mã giao dịch).",
                "Hóa đơn & Trả bàn: Tự động sinh Invoice (mã UUID), tải hóa đơn text (invoice.txt), đơn chuyển COMPLETED và mở lại bàn AVAILABLE.",
            ],
            "Phân hệ Thanh toán do bạn Quỳnh thực hiện là bước cuối cùng trong chu trình. Điểm nổi bật nhất là cơ chế Chống thanh toán trùng: Thu ngân không thể bấm thu tiền hai lần cho cùng một đơn hàng. Sau khi thanh toán thành công, hệ thống tự động xuất hóa đơn có thể tải về máy, đơn hàng chuyển sang COMPLETED và bàn ăn được giải phóng ngay lập tức cho lượt khách tiếp theo."
        ),
        (
            "Slide 9: Thiết Kế Giao Diện Hiện Đại & Tối Giản (Clean UI)",
            [
                "Bảng màu cao cấp: Tông Midnight Slate Navy (#0f172a) và Royal Blue (#2563eb), loại bỏ hoàn toàn màu tím và nút xanh lá cây cũ.",
                "Bố cục trực quan: Sidebar điều hướng thông minh hiển thị số bàn trống, số đơn đang nấu và số món trong giỏ hàng.",
                "Thẻ món & sơ đồ bàn: Sử dụng Grid Card hiện đại, viền mỏng (#e2e8f0), hiệu ứng đổ bóng mượt mà, pill badge trạng thái tinh tế.",
            ],
            "Về mặt trải nghiệm người dùng, em đã nâng cấp toàn diện giao diện Streamlit theo phong cách Clean & Professional. Chúng em sử dụng tông màu Slate Navy kết hợp Royal Blue, tối ưu bố cục dạng lưới thẻ (Grid Cards), thay thế các nút xanh lá cũ bằng nút gradient hiện đại. Giao diện được tối ưu để nhân viên phục vụ, đầu bếp và thu ngân có thể thao tác nhanh chóng trên cả màn hình máy tính lẫn máy tính bảng."
        ),
        (
            "Slide 10: Cấu Trúc Dữ Liệu Tự Cài Đặt & 125 Test Case",
            [
                "Thư mục src/structures/: Cài đặt DoublyLinkedList (O(1) thêm/xóa đầu cuối), HashTable (Separate Chaining), Stack (LIFO hoàn tác).",
                "Chất lượng kiểm thử: 125/125 test case PASSED 100% qua pytest.",
                "Bao phủ toàn diện: test_cart_core, test_order_logic, test_kitchen_prep_time, test_payment_service, test_structures, test_smoke.",
            ],
            "Để đáp ứng yêu cầu học thuật khắt khe của môn Cấu trúc dữ liệu và giải thuật, bạn Lan Anh đã tự tay cài đặt các lớp Danh sách liên kết đôi, Bảng băm và Ngăn xếp thay vì chỉ dùng thư viện có sẵn. Đồng thời, cả nhóm đã xây dựng bộ kiểm thử tự động gồm 125 test case. Trước khi hợp nhất vào nhánh dev, mọi dòng code đều phải vượt qua 100% test case, đảm bảo hệ thống tuyệt đối không phát sinh lỗi tiềm ẩn."
        ),
        (
            "Slide 11: Minh Họa Luồng Nghiệp Vụ Thực Tế (Demo)",
            [
                "demo_mentor_flow.py: Mô phỏng toàn bộ luồng từ Bàn trống -> Đặt món -> Bếp nấu -> Thanh toán -> Xuất hóa đơn -> Giải phóng bàn.",
                "demo_order.py: Thử nghiệm sửa món, hủy món, áp mã giảm giá, kiểm tra chặn thanh toán trùng.",
                "demo_data_structure.py: Trực quan hóa hàng đợi bếp FIFO và cơ chế Undo/Redo giỏ hàng.",
                "Web Streamlit: Trải nghiệm ứng dụng thực tế tại localhost:8501.",
            ],
            "Để chứng minh tính thông suốt của phần mềm, nhóm đã chuẩn bị 3 script demo độc lập có thể chạy trực tiếp qua Terminal cũng như hệ thống Web Streamlit. Script demo_mentor_flow.py mô phỏng chính xác kịch bản nghiệm thu mà Mentor mong muốn: từ một chiếc bàn trống, khách gọi món, bếp nấu, tính tiền, in hóa đơn và bàn tự động sẵn sàng đón khách mới."
        ),
        (
            "Slide 12: Tổng Kết, Bài Học Kinh Nghiệm & Hỏi Đáp",
            [
                "Kết quả: Tiếp thu 100% góp ý mentor; hoàn thành 8/8 task trên Google Sheets; 125 test pass; giao diện Clean hiện đại.",
                "Bài học: Phân chia kiến trúc rõ ràng từ đầu; kỹ năng làm việc nhóm trên Git qua nhánh dev; tự động hóa kiểm thử phần mềm.",
                "Lời cảm ơn: Nhóm xin chân thành cảm ơn Mentor đã đồng hành và định hướng!",
            ],
            "Kính thưa Thầy/Cô, qua dự án này, nhóm không chỉ hoàn thành một phần mềm quản lý F&B hoàn chỉnh mà quan trọng hơn là học được kỹ năng làm việc nhóm thực chiến trên Git, thiết kế kiến trúc chuẩn và tư duy viết mã có kiểm thử bảo vệ. Chúng em xin chân thành cảm ơn những định hướng quý báu của Mentor trong suốt quá trình thực hiện. Sau đây, nhóm em xin phép được trình chiếu demo trực tiếp và lắng nghe các câu hỏi đóng góp từ Thầy/Cô!"
        ),
    ]

    for title, bullets, script in slides_data:
        h = doc.add_heading(title, level=1)
        h.runs[0].font.color.rgb = COLOR_BLUE
        h.runs[0].font.size = Pt(13)
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(4)

        p_b_title = doc.add_paragraph()
        r_b_title = p_b_title.add_run("📌 Nội dung hiển thị trên Slide:")
        r_b_title.font.bold = True
        r_b_title.font.size = Pt(10.5)
        for b in bullets:
            bp = doc.add_paragraph(style='List Bullet')
            br = bp.add_run(b)
            br.font.size = Pt(10)
            bp.paragraph_format.space_after = Pt(2)

        p_s_title = doc.add_paragraph()
        p_s_title.paragraph_format.space_before = Pt(4)
        r_s_title = p_s_title.add_run("🎤 Lời thoại kịch bản thuyết trình:")
        r_s_title.font.bold = True
        r_s_title.font.color.rgb = RGBColor(180, 83, 9)
        r_s_title.font.size = Pt(10.5)

        sp = doc.add_paragraph()
        sr = sp.add_run(f'"{script}"')
        sr.font.italic = True
        sr.font.size = Pt(10.5)
        sr.font.color.rgb = RGBColor(15, 23, 42)
        sp.paragraph_format.space_after = Pt(10)

    out_path = "Slide_Va_Kich_Ban_Bao_Cao_NguyenLeVy.docx"
    doc.save(out_path)
    print("SUCCESS: Da tao thanh cong file Word Slide_Va_Kich_Ban_Bao_Cao_NguyenLeVy.docx")

if __name__ == "__main__":
    generate_slide_docx()
